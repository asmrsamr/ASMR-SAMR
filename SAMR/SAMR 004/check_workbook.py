"""
SAMR Production Suite — acceptance checks.

Run:  python check_workbook.py

Works on a COPY of the workbook so the delivered file stays at its defaults.
Each scenario writes input cells with openpyxl, recalculates with Excel
(tools/excel_recalc.ps1), then reads the computed values back and asserts.
Expected values are computed from the SAME constants the builder uses
(build_workbook.FM_ROWS / INV_ROWS), so formula tweaks keep the checks honest.
"""

import datetime
import os
import shutil
import subprocess
import sys

from openpyxl import load_workbook

from build_workbook import FM_FIRST, FM_ROWS, INV_FIRST, INV_ROWS

sys.stdout.reconfigure(encoding="utf-8", errors="replace")

HERE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(HERE, "SAMR_Production_Suite.xlsx")
TEST = os.path.join(HERE, "_acceptance_test_copy.xlsx")
RECALC = os.path.join(HERE, "tools", "excel_recalc.ps1")

ERRORS = ("#REF!", "#NAME?", "#DIV/0!", "#VALUE!", "#N/A", "#NULL!", "#NUM!")

FM = "Formula Master"
BC = "Batch Calculator"
INV = "Inventory"
LOG = "Production Log"
YC = "Yield & Capacity"
PR = "Pricing & Sales"

DPG_NAME = "DPG (Dipropylene Glycol)"
ALC_NAME = "Alcohol 190 proof SDA 40B"
SHIPPING = 351.0

# ---------------- expected values from the builder's own constants ----------
_inv = {row[0]: row for row in INV_ROWS}


def pack_grams(mat):
    _, qty, unit, _price, dens, *_ = _inv[mat]
    return qty * dens if unit == "ml" else qty


def cost_per_g(mat):
    return _inv[mat][3] / pack_grams(mat)


DPG_CPG = cost_per_g(DPG_NAME)


def weighed_cpg(mat, d10):
    return 0.1 * cost_per_g(mat) + 0.9 * DPG_CPG if d10 else cost_per_g(mat)


CONC_CPG = sum(pct * weighed_cpg(mat, d10) for _ph, mat, pct, d10, _o, _n in FM_ROWS)
MATERIALS_TOTAL = sum(row[3] for row in INV_ROWS)          # 1 pack of everything
SHIP_LOAD = 1 + SHIPPING / MATERIALS_TOTAL
CONC_LOADED = CONC_CPG * SHIP_LOAD
ALC_LOADED = cost_per_g(ALC_NAME) * SHIP_LOAD
JUICE_50 = 50 * 0.85 * 0.28 * CONC_LOADED + 50 * 0.85 * 0.72 * ALC_LOADED

OPT_SUM = sum(pct for _p, _m, pct, _d, opt, _n in FM_ROWS if opt)
D10_ACTIVE = sum(pct for _p, _m, pct, d10, _o, _n in FM_ROWS if d10)

CAPS = {mat: pack_grams(mat) / (pct * (0.1 if d10 else 1.0))
        for _p, mat, pct, d10, _o, _n in FM_ROWS}
LIMIT_MAT = min(CAPS, key=CAPS.get)
LIMIT_G = CAPS[LIMIT_MAT]


def bottles(size_ml):
    return int(LIMIT_G // (size_ml * 0.85 * 0.28))


def fm_row(mat):
    return FM_FIRST + [r[1] for r in FM_ROWS].index(mat)


def inv_row(mat):
    return INV_FIRST + [r[0] for r in INV_ROWS].index(mat)


def yc_row(mat):
    return 27 + [r[1] for r in FM_ROWS].index(mat)


HED = "Hedione (Firmenich)"
HED_PCT = dict((r[1], r[2]) for r in FM_ROWS)[HED]

failures = []
passed = 0


def check(label, actual, expected, tol=0.005):
    global passed
    try:
        ok = abs(float(actual) - float(expected)) <= tol
    except (TypeError, ValueError):
        ok = False
    if ok:
        passed += 1
        print(f"  PASS  {label}: {actual}")
    else:
        failures.append(label)
        print(f"  FAIL  {label}: got {actual!r}, expected {expected} (±{tol})")


def recalc(path):
    r = subprocess.run(
        ["powershell.exe", "-NoProfile", "-ExecutionPolicy", "Bypass",
         "-File", RECALC, "-Path", path],
        capture_output=True, text=True, timeout=300)
    if "RECALC_OK" not in r.stdout:
        print(r.stdout, r.stderr)
        sys.exit("Excel recalculation failed — is Excel installed and closed?")


def set_cells(path, cells):
    wb = load_workbook(path)
    for (sheet, addr), value in cells.items():
        wb[sheet][addr] = value
    wb.calculation.fullCalcOnLoad = True
    wb.save(path)
    wb.close()


def values(path):
    return load_workbook(path, data_only=True)


def scan_errors(wb, label):
    global passed
    found = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value in ERRORS:
                    found.append(f"{ws.title}!{c.coordinate} = {c.value}")
    if found:
        failures.append(f"formula errors ({label})")
        print(f"  FAIL  formula errors after {label}:")
        for f in found[:20]:
            print(f"        {f}")
    else:
        passed += 1
        print(f"  PASS  no formula errors ({label})")


def main():
    if not os.path.exists(SRC):
        sys.exit(f"{SRC} not found — run build_workbook.py first")
    shutil.copyfile(SRC, TEST)

    # ---- scenario 1: defaults (Extrait, 50 ml, 28%, 0.85) ----------------
    print("\n[1] Defaults — Formula Master sums + Extrait 50 ml @ 28%")
    recalc(TEST)
    wb = values(TEST)
    check("input % sum = 100%", wb[FM]["C33"].value, 1.0, 1e-9)
    check("active % sum = 100% (optionals ON)", wb[FM]["F33"].value, 1.0, 1e-9)
    check("default format is Extrait", 1.0 if wb[BC]["C4"].value == "Extrait de Parfum" else 0.0, 1.0)
    check("total batch mass = 42.5 g", wb[BC]["C21"].value, 42.5)
    check("concentrate = 11.9 g (28%)", wb[BC]["C23"].value, 11.9)
    check("per-material sum = concentrate (±0.01)", wb[BC]["E54"].value, 11.9, 0.01)
    check("alcohol = 30.6 g", wb[BC]["D58"].value, 30.6)
    check("batch total = 42.5 g", wb[BC]["D69"].value, 42.5)
    scan_errors(wb, "defaults")
    wb.close()

    # ---- scenario 2: optionals OFF → renormalized to 100% ----------------
    print("\n[2] Optionals OFF — renormalization")
    set_cells(TEST, {(FM, "C4"): "NO"})
    recalc(TEST)
    wb = values(TEST)
    check("active % sum = 100% (optionals OFF)", wb[FM]["F33"].value, 1.0, 1e-9)
    check("Damascone Beta active = 0", wb[FM][f"F{fm_row('Damascone Beta')}"].value, 0.0, 1e-12)
    check("Saffron Accord active = 0", wb[FM][f"F{fm_row('Saffron Accord')}"].value, 0.0, 1e-12)
    check("Hedione renormalized", wb[FM][f"F{fm_row(HED)}"].value, HED_PCT / (1 - OPT_SUM), 1e-9)
    check("weigh sum still = concentrate", wb[BC]["E54"].value, 11.9, 0.01)
    scan_errors(wb, "optionals OFF")
    wb.close()
    set_cells(TEST, {(FM, "C4"): "YES"})

    # ---- scenario 2b: Yield & Capacity (fresh inventory) -----------------
    print("\n[2b] Yield & Capacity — bottleneck + bottle counts")
    recalc(TEST)
    wb = values(TEST)
    check(f"max concentrate = {LIMIT_G:.2f} g ({LIMIT_MAT} limit)", wb[YC]["B11"].value, LIMIT_G, 0.1)
    check(f"limiting material is {LIMIT_MAT}",
          1.0 if wb[YC]["F11"].value == LIMIT_MAT else 0.0, 1.0)
    check(f"10 ml Extrait bottles = {bottles(10)}", wb[YC]["F16"].value, bottles(10))
    check(f"30 ml Extrait bottles = {bottles(30)}", wb[YC]["F18"].value, bottles(30))
    check(f"50 ml Extrait bottles = {bottles(50)}", wb[YC]["F19"].value, bottles(50))
    check(f"100 ml Extrait bottles = {bottles(100)}", wb[YC]["F21"].value, bottles(100))
    check("Bergamot Accord capacity row", wb[YC][f"F{yc_row('Bergamot Accord')}"].value,
          CAPS["Bergamot Accord"], 0.1)
    check("Iso E Super capacity row", wb[YC][f"F{yc_row('Iso E Super')}"].value,
          CAPS["Iso E Super"], 0.1)
    scan_errors(wb, "yield & capacity")
    wb.close()

    # ---- scenario 2c: Pricing & Sales -----------------------------------
    print("\n[2c] Pricing & Sales — cost → retail")
    recalc(TEST)
    wb = values(TEST)
    check(f"50 ml juice cost ≈ {JUICE_50:.2f} SAR", wb[PR]["B13"].value, JUICE_50, 0.2)
    check("50 ml total cost = juice + 40 pkg", wb[PR]["E13"].value,
          round(wb[PR]["B13"].value + 40, 2), 0.01)
    check("50 ml retail ex-VAT = cost × 3", wb[PR]["F13"].value, wb[PR]["E13"].value * 3, 0.01)
    check("50 ml retail inc-VAT = ex × 1.15", wb[PR]["G13"].value, wb[PR]["F13"].value * 1.15, 0.01)
    check("batch profit = 4 × (retail − cost)", wb[PR]["C24"].value,
          4 * (wb[PR]["F13"].value - wb[PR]["E13"].value), 0.01)
    scan_errors(wb, "pricing & sales")
    wb.close()

    # ---- scenario 3: water-based mist 100 g @ 2.5%, 4x -------------------
    print("\n[3] Water-Based Mist — 100 g @ 2.5%, solubilizer 4×")
    set_cells(TEST, {(BC, "C4"): "Water-Based Mist", (BC, "C5"): 100,
                     (BC, "C6"): "g", (BC, "C7"): 0.025})
    recalc(TEST)
    wb = values(TEST)
    check("fragrance = 2.5 g", wb[BC]["C23"].value, 2.5)
    check("Polysorbate 20 = 10 g", wb[BC]["D61"].value, 10.0)
    check("glycerin = 2.5 g", wb[BC]["D60"].value, 2.5)
    check("preservative = 0.8 g", wb[BC]["D62"].value, 0.8)
    check("water = 84.2 g", wb[BC]["D59"].value, 84.2)
    check("total = 100 g", wb[BC]["D69"].value, 100.0)
    scan_errors(wb, "water mist")
    wb.close()

    # ---- scenario 4: cream route 2, 100 g @ 1% ---------------------------
    print("\n[4] Cream Route 2 — 100 g @ 1%")
    set_cells(TEST, {(BC, "C4"): "Cream Route 2 (from scratch)", (BC, "C7"): 0.01})
    recalc(TEST)
    wb = values(TEST)
    check("fragrance = 1 g", wb[BC]["C23"].value, 1.0)
    check("water = 67.5 g", wb[BC]["D59"].value, 67.5)
    check("e-wax = 5 g", wb[BC]["D64"].value, 5.0)
    check("carrier oil = 15 g", wb[BC]["D66"].value, 15.0)
    check("total = exactly 100 g", wb[BC]["D69"].value, 100.0, 1e-9)
    scan_errors(wb, "cream route 2")
    wb.close()

    # ---- scenario 5: spec EDP verification + batch depletes inventory ----
    print("\n[5] EDP 50 ml @ 20% verification + Production Log depletes Inventory")
    today = datetime.date.today()
    set_cells(TEST, {
        (BC, "C4"): "Eau de Parfum", (BC, "C5"): 50, (BC, "C6"): "ml", (BC, "C7"): 0.20,
        (LOG, "A5"): today, (LOG, "B5"): "TEST-01", (LOG, "C5"): "Eau de Parfum",
        (LOG, "D5"): 50, (LOG, "E5"): "ml", (LOG, "F5"): 0.20,
        (LOG, "G5"): 8.5, (LOG, "H5"): 34.0, (LOG, "I5"): today, (LOG, "K5"): "Macerating",
    })
    recalc(TEST)
    wb = values(TEST)
    check("EDP total mass = 42.5 g", wb[BC]["C21"].value, 42.5)
    check("EDP concentrate = 8.5 g", wb[BC]["C23"].value, 8.5)
    check("EDP per-material sum (±0.01)", wb[BC]["E54"].value, 8.5, 0.01)
    check("EDP alcohol = 34 g", wb[BC]["D58"].value, 34.0)
    hed_r = inv_row(HED)
    mal_r = inv_row("Maltol (Crystals)")
    dpg_r = inv_row(DPG_NAME)
    alc_r = inv_row(ALC_NAME)
    check(f"Hedione used = 8.5×{HED_PCT:.1%}", wb[INV][f"J{hed_r}"].value, 8.5 * HED_PCT, 1e-3)
    check("Hedione on hand = 27 − used", wb[INV][f"K{hed_r}"].value, 27 - 8.5 * HED_PCT, 1e-3)
    check("Maltol (D10) used neat = 8.5×2%×0.1", wb[INV][f"J{mal_r}"].value, 8.5 * 0.02 * 0.1, 1e-4)
    check(f"DPG used = 8.5×{D10_ACTIVE:.1%}×0.9", wb[INV][f"J{dpg_r}"].value,
          8.5 * D10_ACTIVE * 0.9, 1e-4)
    check("alcohol used = 34 g", wb[INV][f"J{alc_r}"].value, 34.0)
    check("alcohol on hand = 4050 − 34", wb[INV][f"K{alc_r}"].value, 4050 - 34.0)
    scan_errors(wb, "logged batch")
    wb.close()

    # ---- scenario 6: heavy usage triggers low-stock alerts ----------------
    print("\n[6] Low-stock path — big batch forces alerts on the Dashboard")
    set_cells(TEST, {(LOG, "G6"): 100.0, (LOG, "B6"): "TEST-02"})
    recalc(TEST)
    wb = values(TEST)
    dash = wb["Dashboard"]["G7"].value
    msg = wb["Dashboard"]["B18"].value
    if isinstance(dash, (int, float)) and dash > 0 and isinstance(msg, str) and msg.startswith("⚠"):
        globals()["passed"] += 1
        print(f"  PASS  low-stock alert renders ({int(dash)} items): {msg[:70]}…")
    else:
        failures.append("low-stock alert message")
        print(f"  FAIL  low-stock alert: count={dash!r}, msg={msg!r}")
    scan_errors(wb, "low stock")
    wb.close()

    # ---- scenario 7: every inventory row carries a working product link ---
    print("\n[7] Inventory — product-link column present on every row")
    wb = load_workbook(TEST)
    ws = wb[INV]
    missing = [ws[f"A{r}"].value for r in range(INV_FIRST, INV_FIRST + len(INV_ROWS))
               if not (ws[f"N{r}"].hyperlink
                       and str(ws[f"N{r}"].hyperlink.target).startswith("https://iterji.com/"))]
    if missing:
        failures.append("inventory links")
        print(f"  FAIL  rows without an iterji.com hyperlink: {missing}")
    else:
        globals()["passed"] += 1
        print(f"  PASS  all {len(INV_ROWS)} inventory rows link to iterji.com")
    wb.close()

    # ---- summary ----------------------------------------------------------
    os.remove(TEST)
    print(f"\n{'=' * 60}\n{passed} checks passed, {len(failures)} failed")
    if failures:
        for f in failures:
            print(f"  FAILED: {f}")
        sys.exit(1)
    print("ALL ACCEPTANCE TESTS PASSED")


if __name__ == "__main__":
    main()

