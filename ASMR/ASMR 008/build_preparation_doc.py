from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path("ASMR Preparation Guide.docx")

BLUE = RGBColor(46, 116, 181)
DARK = RGBColor(45, 45, 45)
MUTED = RGBColor(90, 90, 90)
FILL = "E8EEF5"


def set_font(run, size=11, bold=False, italic=False, color=DARK, name="Calibri"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run._element.rPr.rFonts.set(qn("w:cs"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = color


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text, bold=False, fill=None):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(str(text))
    set_font(run, size=9.5, bold=bold)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    if fill:
        shade_cell(cell, fill)


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.space_before = Pt(14 if level == 1 else 10)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    set_font(run, size=16 if level == 1 else 13, bold=True, color=BLUE)


def add_para(doc, text, bold=False, italic=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.25
    run = p.add_run(text)
    set_font(run, bold=bold, italic=italic)


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.25
    run = p.add_run(text)
    set_font(run)


def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.25
    run = p.add_run(text)
    set_font(run)


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = False
    for i, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], header, bold=True, fill=FILL)
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], value)
    if widths:
        for row in table.rows:
            for i, width in enumerate(widths):
                row.cells[i].width = Inches(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def build():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("ASMR Preparation Guide")
    set_font(run, size=24, bold=True)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(16)
    run = subtitle.add_run("Personal skin-scent production, scale-up, and selling readiness")
    set_font(run, size=12, italic=True, color=MUTED)

    add_heading(doc, "Project Identity")
    add_para(doc, "ASMR is a personal signature-scent project built around a 37 C skin-temperature idea: transparent citrus, radiant floral heart, and warm musky-woody drydown.")
    add_table(
        doc,
        ["Field", "Current setting"],
        [
            ["Default format", "Eau de Parfum at 20%"],
            ["Default 50 ml finished mass", "42.5 g"],
            ["Default 50 ml EDP concentrate", "8.5 g"],
            ["Current bottleneck", "Iso E Super"],
            ["Production tabs", "Yield & Capacity, Scale-Up Plan, Pricing & Sales, Selling Strategy"],
        ],
        widths=[2.2, 4.1],
    )

    add_heading(doc, "Before You Start")
    add_para(doc, "Prepare the bench before opening materials. Keep the printed Batch Calculator beside the scale and tick every line after weighing.")
    for item in [
        "Use a calibrated 0.01 g scale, clean glassware, pipettes, labels, gloves, and ventilation.",
        "Weigh in grams only. Never use drops.",
        "Use Bergamot Accord for skin products. Bergamot EO is phototoxic and is not part of the perfume formula.",
        "Use one amber material only: Amberxan / Ambroxan in this formula.",
        "Water-containing products require preservative. Alcohol products self-preserve.",
        "Verify against current IFRA standards before commercial sale.",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "D10 Dilution Prep")
    add_para(doc, "D10 means a 10% dilution in DPG. For any D10 material:")
    add_para(doc, "X g of D10 = 0.1X g neat material + 0.9X g DPG", bold=True)
    add_table(
        doc,
        ["Material", "Default D10 batch", "Neat", "DPG"],
        [
            ["Pink Peppercorn", "10.00 g", "1.00 g", "9.00 g"],
            ["Floralozone", "10.00 g", "1.00 g", "9.00 g"],
            ["Irone Alpha", "10.00 g", "1.00 g", "9.00 g"],
            ["Javanol", "10.00 g", "1.00 g", "9.00 g"],
            ["Safraleine", "10.00 g", "1.00 g", "9.00 g"],
        ],
        widths=[2.1, 1.35, 1.2, 1.2],
    )

    add_heading(doc, "Concentrate Build")
    for step in [
        "Open the workbook, set format, size, concentration, density, and optional-material toggle.",
        "Print the Batch Calculator and write the batch ID on both the sheet and the empty concentrate vessel.",
        "Tare the empty vessel to 0.00 g.",
        "Weigh BASE materials first, then HEART, then TOP. Tare between each material and tick the printed line.",
        "For D10 rows, weigh the dilution amount shown by the workbook, not the neat-material amount.",
        "Cap and roll the concentrate for one minute after the last material is added.",
        "Rest the concentrate 24-48 hours before alcohol dilution.",
        "Label the concentrate with formula, batch ID, date, optional status, and total concentrate grams.",
    ]:
        add_number(doc, step)
    add_table(
        doc,
        ["Order", "ASMR phase", "Handling note"],
        [
            ["1. BASE", "Iso E Super, Amberxan, musks, sandalwoods, cedar, vetiver, Safraleine", "This controls longevity and the skin-like aura."],
            ["2. HEART", "Hedione, orange blossom, ionones, benzyl salicylate, linalool", "This gives diffusion and floral warmth."],
            ["3. TOP", "Bergamot accord, petitgrain, neroli, pink pepper, Floralozone", "Volatile materials go last to reduce loss to air."],
        ],
        widths=[1.0, 3.4, 1.9],
    )

    add_heading(doc, "Dilution, Maceration, and Bottling")
    for step in [
        "For a 50 ml EDP at 20%, use 8.5 g concentrate and 34.0 g alcohol at density 0.85 g/ml.",
        "Add alcohol slowly, cap, and roll until the blend looks uniform.",
        "Log the batch immediately in Production Log so inventory, yield, scale-up, and selling tabs update.",
        "Macerate 4-8 weeks in a dark, cool place. Two weeks is minimum; six weeks is the standard target.",
        "Swirl once daily during week one, then let the bottle rest.",
        "Cold crash and filter if haze or sediment appears.",
        "Bottle by weight. For 50 ml at 0.85 g/ml, fill to 42.5 g finished perfume.",
        "Keep one retained sample and a complete batch record for every batch sold.",
    ]:
        add_number(doc, step)

    add_heading(doc, "Scale-Up and Selling")
    add_para(doc, "The workbook now includes Scale-Up Plan and Selling Strategy tabs. Change inventory pack count, pack size, or density and the bottle counts update dynamically.")
    for item in [
        "Inventory includes product links and a per-material 50 ml bottle count.",
        "Yield & Capacity includes fixed 50 ml bottle count and editable bottle-size count.",
        "Scale-Up Plan highlights ingredients to increase to hit the target production level.",
        "Selling Strategy proposes pricing, channel split, and go/no-go readiness.",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "Safety Checklist")
    for item in [
        "Patch-test every finished product.",
        "Alcohol is flammable: work ventilated, away from flame or heat.",
        "Water-based mist and cream route 2 require preservative.",
        "Leave-on creams are stricter than fine fragrance; keep fragrance at or below 1.2%.",
        "Store raw materials dark, cool, and tightly closed.",
    ]:
        add_bullet(doc, item)

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
