"""
ASMR Production Suite — acceptance checks.

Run:  python check_workbook.py

Works on a COPY of the workbook so the delivered file stays at its defaults.
Each scenario writes input cells with openpyxl, recalculates with Excel
(tools/excel_recalc.ps1), then reads the computed values back and asserts.
"""

import datetime
import os
import shutil
import subprocess
import sys

from openpyxl import load_workbook

sys.stdout.reconfigure(encoding="utf-8", errors="replace")

HERE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(HERE, "ASMR_Production_Suite.xlsx")
TEST = os.path.join(HERE, "_acceptance_test_copy.xlsx")
RECALC = os.path.join(HERE, "tools", "excel_recalc.ps1")

ERRORS = ("#REF!", "#NAME?", "#DIV/0!", "#VALUE!", "#N/A", "#NULL!", "#NUM!")

# key cell addresses (must match build_workbook.py layout)
FM = "Formula Master"
BC = "Batch Calculator"
INV = "Inventory"
LOG = "Production Log"

failures = []
passed = 0


def check(label, actual, expected, tol=0.005):
    global passed
    try:
        ok = abs(float(actual) - float(expected)) <= tol
    except (TypeError, ValueError):
        ok = False
    if ok:
        passed += 1
        print(f"  PASS  {label}: {actual}")
    else:
        failures.append(label)
        print(f"  FAIL  {label}: got {actual!r}, expected {expected} (±{tol})")


def recalc(path):
    r = subprocess.run(
        ["powershell.exe", "-NoProfile", "-ExecutionPolicy", "Bypass",
         "-File", RECALC, "-Path", path],
        capture_output=True, text=True, timeout=300)
    if "RECALC_OK" not in r.stdout:
        print(r.stdout, r.stderr)
        sys.exit("Excel recalculation failed — is Excel installed and closed?")


def set_cells(path, cells):
    wb = load_workbook(path)
    for (sheet, addr), value in cells.items():
        wb[sheet][addr] = value
    wb.calculation.fullCalcOnLoad = True
    wb.save(path)
    wb.close()


def values(path):
    return load_workbook(path, data_only=True)


def scan_errors(wb, label):
    global passed
    found = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value in ERRORS:
                    found.append(f"{ws.title}!{c.coordinate} = {c.value}")
    if found:
        failures.append(f"formula errors ({label})")
        print(f"  FAIL  formula errors after {label}:")
        for f in found[:20]:
            print(f"        {f}")
    else:
        passed += 1
        print(f"  PASS  no formula errors ({label})")


def main():
    if not os.path.exists(SRC):
        sys.exit(f"{SRC} not found — run build_workbook.py first")
    shutil.copyfile(SRC, TEST)

    # ---- scenario 1: defaults (Extrait, 50 ml, 28%, 0.85) ----------------
    print("\n[1] Defaults — Formula Master sums + Extrait 50 ml @ 28%")
    recalc(TEST)
    wb = values(TEST)
    check("input % sum = 100%", wb[FM]["C33"].value, 1.0, 1e-9)
    check("active % sum = 100% (optionals ON)", wb[FM]["F33"].value, 1.0, 1e-9)
    check("default format is Extrait", 1.0 if wb[BC]["C4"].value == "Extrait de Parfum" else 0.0, 1.0)
    check("total batch mass = 42.5 g", wb[BC]["C21"].value, 42.5)
    check("concentrate = 11.9 g (28%)", wb[BC]["C23"].value, 11.9)
    check("per-material sum = concentrate (±0.01)", wb[BC]["E54"].value, 11.9, 0.01)
    check("alcohol = 30.6 g", wb[BC]["D58"].value, 30.6)
    check("batch total = 42.5 g", wb[BC]["D69"].value, 42.5)
    scan_errors(wb, "defaults")
    wb.close()

    # ---- scenario 2: optionals OFF → renormalized to 100% ----------------
    print("\n[2] Optionals OFF — renormalization")
    set_cells(TEST, {(FM, "C4"): "NO"})
    recalc(TEST)
    wb = values(TEST)
    check("active % sum = 100% (optionals OFF)", wb[FM]["F33"].value, 1.0, 1e-9)
    check("Irone Alpha active = 0", wb[FM]["F15"].value, 0.0, 1e-12)
    check("Javanol active = 0", wb[FM]["F23"].value, 0.0, 1e-12)
    check("Hedione renormalized = 20/98.5", wb[FM]["F12"].value, 0.20 / 0.985, 1e-9)
    check("weigh sum still = concentrate", wb[BC]["E54"].value, 11.9, 0.01)
    scan_errors(wb, "optionals OFF")
    wb.close()
    set_cells(TEST, {(FM, "C4"): "YES"})

    # ---- scenario 2b: Yield & Capacity (fresh inventory) -----------------
    print("\n[2b] Yield & Capacity — bottleneck + bottle counts")
    YC = "Yield & Capacity"
    recalc(TEST)
    wb = values(TEST)
    check("max concentrate = 60 g (Iso E Super limit)", wb[YC]["B11"].value, 60.0, 0.1)
    check("limiting material is Iso E Super",
          1.0 if wb[YC]["F11"].value == "Iso E Super" else 0.0, 1.0)
    check("10 ml Extrait bottles = 25", wb[YC]["F16"].value, 25)
    check("30 ml Extrait bottles = 8", wb[YC]["F18"].value, 8)
    check("50 ml Extrait bottles = 5", wb[YC]["F19"].value, 5)
    check("100 ml Extrait bottles = 2", wb[YC]["F21"].value, 2)
    check("Bergamot Accord capacity row = 100 g", wb[YC]["F27"].value, 100.0, 0.1)
    check("Iso E Super capacity row = 60 g", wb[YC]["F38"].value, 60.0, 0.1)
    scan_errors(wb, "yield & capacity")
    wb.close()

    # ---- scenario 2c: Pricing & Sales -----------------------------------
    print("\n[2c] Pricing & Sales — cost → retail")
    PR = "Pricing & Sales"
    recalc(TEST)
    wb = values(TEST)
    check("50 ml juice cost ≈ 74.82 SAR", wb[PR]["B15"].value, 74.82, 0.2)
    check("50 ml total cost = juice + 40 pkg", wb[PR]["E15"].value,
          round(wb[PR]["B15"].value + 40, 2), 0.01)
    check("50 ml cost per ml = total / 50", wb[PR]["F15"].value, wb[PR]["E15"].value / 50, 0.01)
    check("50 ml retail ex-VAT = cost × 3", wb[PR]["G15"].value, wb[PR]["E15"].value * 3, 0.01)
    check("50 ml retail per ml = retail / 50", wb[PR]["H15"].value, wb[PR]["G15"].value / 50, 0.01)
    check("50 ml retail inc-VAT = ex × 1.15", wb[PR]["I15"].value, wb[PR]["G15"].value * 1.15, 0.01)
    check("batch profit = 4 × (retail − cost)", wb[PR]["C27"].value,
          4 * (wb[PR]["G15"].value - wb[PR]["E15"].value), 0.01)
    scan_errors(wb, "pricing & sales")
    wb.close()

    # ---- scenario 3: water-based mist 100 g @ 2.5%, 4x -------------------
    print("\n[3] Water-Based Mist — 100 g @ 2.5%, solubilizer 4×")
    set_cells(TEST, {(BC, "C4"): "Water-Based Mist", (BC, "C5"): 100,
                     (BC, "C6"): "g", (BC, "C7"): 0.025})
    recalc(TEST)
    wb = values(TEST)
    check("fragrance = 2.5 g", wb[BC]["C23"].value, 2.5)
    check("Polysorbate 20 = 10 g", wb[BC]["D61"].value, 10.0)
    check("glycerin = 2.5 g", wb[BC]["D60"].value, 2.5)
    check("preservative = 0.8 g", wb[BC]["D62"].value, 0.8)
    check("water = 84.2 g", wb[BC]["D59"].value, 84.2)
    check("total = 100 g", wb[BC]["D69"].value, 100.0)
    scan_errors(wb, "water mist")
    wb.close()

    # ---- scenario 4: cream route 2, 100 g @ 1% ---------------------------
    print("\n[4] Cream Route 2 — 100 g @ 1%")
    set_cells(TEST, {(BC, "C4"): "Cream Route 2 (from scratch)", (BC, "C7"): 0.01})
    recalc(TEST)
    wb = values(TEST)
    check("fragrance = 1 g", wb[BC]["C23"].value, 1.0)
    check("water = 67.5 g", wb[BC]["D59"].value, 67.5)
    check("e-wax = 5 g", wb[BC]["D64"].value, 5.0)
    check("carrier oil = 15 g", wb[BC]["D66"].value, 15.0)
    check("total = exactly 100 g", wb[BC]["D69"].value, 100.0, 1e-9)
    scan_errors(wb, "cream route 2")
    wb.close()

    # ---- scenario 5: spec EDP verification + batch depletes inventory ----
    print("\n[5] EDP 50 ml @ 20% verification + Production Log depletes Inventory")
    today = datetime.date.today()
    set_cells(TEST, {
        (BC, "C4"): "Eau de Parfum", (BC, "C5"): 50, (BC, "C6"): "ml", (BC, "C7"): 0.20,
        (LOG, "A5"): today, (LOG, "B5"): "TEST-01", (LOG, "C5"): "Eau de Parfum",
        (LOG, "D5"): 50, (LOG, "E5"): "ml", (LOG, "F5"): 0.20,
        (LOG, "G5"): 8.5, (LOG, "H5"): 34.0, (LOG, "I5"): today, (LOG, "K5"): "Macerating",
    })
    recalc(TEST)
    wb = values(TEST)
    check("EDP total mass = 42.5 g", wb[BC]["C21"].value, 42.5)
    check("EDP concentrate = 8.5 g", wb[BC]["C23"].value, 8.5)
    check("EDP per-material sum (±0.01)", wb[BC]["E54"].value, 8.5, 0.01)
    check("EDP alcohol = 34 g", wb[BC]["D58"].value, 34.0)
    # Inventory rows: 4=Hedione, 29=Safraleine, 30=DPG, 31=Alcohol
    check("Hedione used = 8.5×20% = 1.70 g", wb[INV]["J4"].value, 1.70)
    check("Hedione on hand = 27−1.7 = 25.30 g", wb[INV]["K4"].value, 25.30)
    check("Safraleine (D10) used neat = 0.0085 g", wb[INV]["J29"].value, 0.0085, 1e-4)
    check("DPG used = 8.5×3.5%×0.9 = 0.268 g", wb[INV]["J30"].value, 0.26775, 1e-4)
    check("alcohol used = 34 g", wb[INV]["J31"].value, 34.0)
    check("alcohol on hand = 810−34 = 776 g", wb[INV]["K31"].value, 776.0)
    scan_errors(wb, "logged batch")
    wb.close()

    # ---- scenario 6: heavy usage triggers low-stock alerts ----------------
    print("\n[6] Low-stock path — big batch forces alerts on the Dashboard")
    set_cells(TEST, {(LOG, "G6"): 60.0, (LOG, "B6"): "TEST-02"})
    recalc(TEST)
    wb = values(TEST)
    dash = wb["Dashboard"]["G7"].value
    msg = wb["Dashboard"]["B18"].value
    if isinstance(dash, (int, float)) and dash > 0 and isinstance(msg, str) and msg.startswith("⚠"):
        globals()["passed"] += 1
        print(f"  PASS  low-stock alert renders ({int(dash)} items): {msg[:70]}…")
    else:
        failures.append("low-stock alert message")
        print(f"  FAIL  low-stock alert: count={dash!r}, msg={msg!r}")
    scan_errors(wb, "low stock")
    wb.close()

    # ---- summary ----------------------------------------------------------
    os.remove(TEST)
    print(f"\n{'=' * 60}\n{passed} checks passed, {len(failures)} failed")
    if failures:
        for f in failures:
            print(f"  FAILED: {f}")
        sys.exit(1)
    print("ALL ACCEPTANCE TESTS PASSED")


if __name__ == "__main__":
    main()
