#!/usr/bin/env python3
"""Extract ASMR & SAMR launch prices/costs from the latest workbook revisions.

The generated SQL intentionally does not directly mutate product stock. Stock
must be adjusted through the authenticated admin dashboard/RPC so movement
history and audit trails stay intact.
"""

from __future__ import annotations

import argparse
import json
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_JSON = ROOT / "data" / "product-pricing-latest.json"
DEFAULT_SQL = ROOT / "supabase" / "imports" / "product-pricing-latest.sql"
EFFECTIVE_DATE = date(2026, 7, 18)
VAT_RATE = 0.15


@dataclass(frozen=True)
class WorkbookSource:
    key: str
    product_id: str
    path: Path
    source_kind: str
    default_price: float | None = None


SOURCES = [
    WorkbookSource(
        "asmr_perfume_008",
        "asmr-extrait",
        ROOT / "ASMR" / "ASMR 008" / "ASMR Perfume 008.xlsx",
        "perfume",
    ),
    WorkbookSource(
        "samr_perfume_011",
        "samr-extrait",
        ROOT / "SAMR" / "SAMR 011" / "SAMR Perfume 011.xlsx",
        "perfume",
    ),
    WorkbookSource(
        "asmr_body_spray_007",
        "asmr-spray",
        ROOT / "ASMR" / "Perfume 007" / "ASMR Body Spray 007.xlsx",
        "body_spray",
        99,
    ),
    WorkbookSource(
        "samr_body_spray_008",
        "samr-spray",
        ROOT / "SAMR" / "SAMR 008" / "SAMR Body Spray 008.xlsx",
        "body_spray",
        79,
    ),
    WorkbookSource(
        "asmr_cream_007",
        "asmr-cream",
        ROOT / "ASMR" / "Perfume 007" / "ASMR Cream 007.xlsx",
        "cream",
        79,
    ),
    WorkbookSource(
        "samr_cream_008",
        "samr-cream",
        ROOT / "SAMR" / "SAMR 008" / "SAMR Cream 008.xlsx",
        "cream",
        69,
    ),
]

PRODUCT_NAMES = {
    "asmr-extrait": "ASMR Extrait de Parfum",
    "samr-extrait": "SAMR Extrait de Parfum",
    "asmr-spray": "ASMR Body Spray",
    "samr-spray": "SAMR Body Spray",
    "asmr-cream": "ASMR Body Cream",
    "samr-cream": "SAMR Body Cream",
    "discovery-set": "Discovery Set",
    "duo-box": "His & Hers Duo Box",
    "asmr-trio": "ASMR Trio Set",
    "samr-trio": "SAMR Trio Set",
}

DEFAULT_SIZES = {
    "asmr-extrait": "50 ml",
    "samr-extrait": "50 ml",
    "asmr-spray": "100 ml",
    "samr-spray": "100 ml",
    "asmr-cream": "100 g",
    "samr-cream": "100 g",
    "discovery-set": "2x 2 ml",
    "duo-box": "Duo Box (2x 50 ml)",
    "asmr-trio": "Trio Set",
    "samr-trio": "Trio Set",
}


def number(value: Any, default: float = 0.0) -> float:
    if value is None or value == "":
        return default
    try:
        return float(value)
    except (TypeError, ValueError):
        return default


def money(value: Any) -> float:
    return round(number(value), 2)


def sql(value: Any) -> str:
    if value is None:
        return "null"
    if isinstance(value, (int, float)):
        return str(round(float(value), 6)).rstrip("0").rstrip(".")
    text = str(value).replace("'", "''")
    return f"'{text}'"


def workbook_modified(path: Path) -> str:
    return datetime.fromtimestamp(path.stat().st_mtime).isoformat(timespec="seconds")


def extract_perfume(source: WorkbookSource) -> dict[str, Any]:
    workbook = load_workbook(source.path, data_only=True, read_only=True)
    selling = workbook["Selling Strategy"]
    pricing = workbook["Pricing & Sales"]
    capacity = workbook["Yield & Capacity"]

    ladder: dict[str, dict[str, float]] = {}
    for row in selling.iter_rows(min_row=12, max_row=17, min_col=1, max_col=9, values_only=True):
        size_raw = row[0]
        if not size_raw:
            continue
        size = str(size_raw).strip()
        if size == "20 ml":
            continue
        ladder[size] = {
            "total_cost": money(row[1]),
            "cost_per_ml": round(number(row[2]), 4),
            "base_retail": money(row[3]),
            "launch_price": money(row[5]),
            "inc_vat": money(row[6]),
            "profit": money(row[7]),
            "margin": round(number(row[8]), 4),
        }

    components: dict[str, dict[str, float]] = {}
    for row in pricing.iter_rows(min_row=12, max_row=17, min_col=1, max_col=11, values_only=True):
        raw_size = row[0]
        if raw_size is None:
            continue
        size = f"{int(raw_size)} ml" if isinstance(raw_size, (int, float)) else str(raw_size).strip()
        if size == "20 ml":
            continue
        components[size] = {
            "ingredient_cost": money(row[1]),
            "packaging_cost": money(row[2]),
            "labor_cost": money(row[3]),
            "total_cost": money(row[4]),
        }

    capacity_by_size: dict[str, dict[str, Any]] = {}
    for row in capacity.iter_rows(min_row=16, max_row=21, min_col=1, max_col=9, values_only=True):
        raw_size = row[0]
        if raw_size is None:
            continue
        size = f"{int(raw_size)} ml" if isinstance(raw_size, (int, float)) else str(raw_size).strip()
        if size == "20 ml":
            continue
        capacity_by_size[size] = {
            "bottles_by_concentrate": round(number(row[3]), 3),
            "bottles_by_alcohol": round(number(row[4]), 3),
            "inventory_capacity_units": int(number(row[5])),
            "cost_per_bottle": money(row[6]),
            "limiting_factor": row[8],
        }

    return {
        "product_id": source.product_id,
        "name": PRODUCT_NAMES[source.product_id],
        "source": str(source.path.relative_to(ROOT)).replace("\\", "/"),
        "source_modified_at": workbook_modified(source.path),
        "source_kind": source.source_kind,
        "default_size": DEFAULT_SIZES[source.product_id],
        "prices": {size: data["launch_price"] for size, data in ladder.items()},
        "costs": components,
        "capacity": capacity_by_size,
        "notes": "Prices use the Selling Strategy launch-price column; costs use Pricing & Sales.",
    }


def extract_body_product(source: WorkbookSource) -> dict[str, Any]:
    workbook = load_workbook(source.path, data_only=True, read_only=True)
    recipe = workbook["Recipe"]
    values = {str(recipe.cell(row, 1).value or "").strip(): number(recipe.cell(row, 2).value) for row in range(1, 28)}

    batch_size = values.get("Batch size (g)", 100)
    concentrate_cost = values.get("Concentrate cost/g", 0)

    if source.source_kind == "body_spray":
        fragrance_pct = values.get("Fragrance %", 0.05)
        alcohol_pct = values.get("Alcohol %", 0.95)
        alcohol_cost = values.get("Alcohol cost/g", 0)
        ingredient_cost = batch_size * fragrance_pct * concentrate_cost + batch_size * alcohol_pct * alcohol_cost
        size = "100 ml"
        retail_note = "Retail price retained from storefront; workbook supplies material cost assumptions only."
    else:
        fragrance_pct = values.get("Fragrance %", 0.01)
        water_pct = 1 - fragrance_pct - 0.05 - 0.15 - 0.008
        ingredient_cost = (
            batch_size * fragrance_pct * concentrate_cost
            + batch_size * water_pct * values.get("Water cost/g", 0)
            + batch_size * 0.05 * values.get("Emulsifying wax cost/g", 0)
            + batch_size * 0.15 * values.get("Carrier oil cost/g", 0)
            + batch_size * 0.008 * values.get("Preservative cost/g", 0)
        )
        size = "100 g"
        retail_note = "Retail price retained from storefront; workbook supplies material cost assumptions only."

    return {
        "product_id": source.product_id,
        "name": PRODUCT_NAMES[source.product_id],
        "source": str(source.path.relative_to(ROOT)).replace("\\", "/"),
        "source_modified_at": workbook_modified(source.path),
        "source_kind": source.source_kind,
        "default_size": size,
        "prices": {size: money(source.default_price)},
        "costs": {
            size: {
                "ingredient_cost": money(ingredient_cost),
                "packaging_cost": 0,
                "labor_cost": 0,
                "total_cost": money(ingredient_cost),
            }
        },
        "capacity": {},
        "notes": retail_note,
    }


def build_manifest() -> dict[str, Any]:
    records = []
    by_id: dict[str, dict[str, Any]] = {}
    for source in SOURCES:
        record = extract_perfume(source) if source.source_kind == "perfume" else extract_body_product(source)
        records.append(record)
        by_id[record["product_id"]] = record

    asmr_50 = by_id["asmr-extrait"]["costs"]["50 ml"]["total_cost"]
    samr_50 = by_id["samr-extrait"]["costs"]["50 ml"]["total_cost"]
    asmr_spray = by_id["asmr-spray"]["costs"]["100 ml"]["total_cost"]
    samr_spray = by_id["samr-spray"]["costs"]["100 ml"]["total_cost"]
    asmr_cream = by_id["asmr-cream"]["costs"]["100 g"]["total_cost"]
    samr_cream = by_id["samr-cream"]["costs"]["100 g"]["total_cost"]

    derived = [
        {
            "product_id": "duo-box",
            "name": PRODUCT_NAMES["duo-box"],
            "source": "Derived from ASMR 50 ml + SAMR 50 ml latest launch prices",
            "source_kind": "derived_set",
            "default_size": DEFAULT_SIZES["duo-box"],
            "prices": {DEFAULT_SIZES["duo-box"]: 499},
            "costs": {
                DEFAULT_SIZES["duo-box"]: {
                    "ingredient_cost": 0,
                    "packaging_cost": 0,
                    "labor_cost": 0,
                    "other_cost": money(asmr_50 + samr_50),
                    "total_cost": money(asmr_50 + samr_50),
                }
            },
            "notes": "Bundle keeps the historical 51 SAR saving: 320 + 230 - 51 = 499. Add coffret packaging cost when supplier quote is known.",
        },
        {
            "product_id": "asmr-trio",
            "name": PRODUCT_NAMES["asmr-trio"],
            "source": "Derived from ASMR 50 ml, ASMR Body Spray, and ASMR Body Cream workbooks",
            "source_kind": "derived_set",
            "default_size": DEFAULT_SIZES["asmr-trio"],
            "prices": {DEFAULT_SIZES["asmr-trio"]: 499},
            "costs": {
                DEFAULT_SIZES["asmr-trio"]: {
                    "ingredient_cost": 0,
                    "packaging_cost": 0,
                    "labor_cost": 0,
                    "other_cost": money(asmr_50 + asmr_spray + asmr_cream),
                    "total_cost": money(asmr_50 + asmr_spray + asmr_cream),
                }
            },
            "notes": "Price matches component launch sum rounded by current storefront strategy.",
        },
        {
            "product_id": "samr-trio",
            "name": PRODUCT_NAMES["samr-trio"],
            "source": "Derived from SAMR 50 ml, SAMR Body Spray, and SAMR Body Cream workbooks",
            "source_kind": "derived_set",
            "default_size": DEFAULT_SIZES["samr-trio"],
            "prices": {DEFAULT_SIZES["samr-trio"]: 379},
            "costs": {
                DEFAULT_SIZES["samr-trio"]: {
                    "ingredient_cost": 0,
                    "packaging_cost": 0,
                    "labor_cost": 0,
                    "other_cost": money(samr_50 + samr_spray + samr_cream),
                    "total_cost": money(samr_50 + samr_spray + samr_cream),
                }
            },
            "notes": "Price matches component launch sum rounded by current storefront strategy.",
        },
        {
            "product_id": "discovery-set",
            "name": PRODUCT_NAMES["discovery-set"],
            "source": "Derived from ASMR/SAMR launch strategy; sample packaging cost still requires supplier quote",
            "source_kind": "derived_set",
            "default_size": DEFAULT_SIZES["discovery-set"],
            "prices": {DEFAULT_SIZES["discovery-set"]: 60},
            "costs": {
                DEFAULT_SIZES["discovery-set"]: {
                    "ingredient_cost": 4,
                    "packaging_cost": 8,
                    "labor_cost": 0,
                    "other_cost": 0,
                    "total_cost": 12,
                }
            },
            "notes": "Sample/vial packaging is estimated and should be replaced with a supplier quote.",
        },
    ]
    records.extend(derived)

    return {
        "generated_at": datetime.now().isoformat(timespec="seconds"),
        "effective_date": EFFECTIVE_DATE.isoformat(),
        "currency": "SAR",
        "vat_rate": VAT_RATE,
        "records": records,
        "stock_policy": "Do not import stock by SQL. Adjust sellable stock in admin so product_stock_movements and audit trail are recorded.",
    }


def variant_rows(manifest: dict[str, Any]) -> list[dict[str, Any]]:
    rows = []
    for record in manifest["records"]:
        for index, (size, price) in enumerate(record["prices"].items(), start=1):
            cost = record["costs"].get(size, {}).get("total_cost", 0)
            variant_name = record["name"] if record["source_kind"] == "derived_set" else f"{record['name']} {size}"
            rows.append(
                {
                    "product_id": record["product_id"],
                    "name": variant_name,
                    "size": size,
                    "price": price,
                    "cost": cost,
                    "tax_rate": VAT_RATE,
                    "is_default": size == record["default_size"],
                    "sort_order": index * 10,
                }
            )
    return rows


def component_rows(manifest: dict[str, Any]) -> list[dict[str, Any]]:
    rows = []
    for record in manifest["records"]:
        size = record["default_size"]
        costs = record["costs"].get(size, {})
        for key, component_type, label in [
            ("ingredient_cost", "ingredient", "Excel launch import: ingredient/juice"),
            ("packaging_cost", "packaging", "Excel launch import: packaging"),
            ("labor_cost", "labor", "Excel launch import: labor"),
            ("other_cost", "other", "Excel launch import: bundle/component cost"),
        ]:
            amount = money(costs.get(key, 0))
            if amount > 0:
                rows.append(
                    {
                        "product_id": record["product_id"],
                        "component_type": component_type,
                        "name": label,
                        "amount": amount,
                        "effective_from": EFFECTIVE_DATE.isoformat(),
                    }
                )
    return rows


def render_sql(manifest: dict[str, Any]) -> str:
    variants = variant_rows(manifest)
    components = component_rows(manifest)
    product_ids = sorted({row["product_id"] for row in variants})
    source_lines = "\n".join(
        f"-- Source: {record['product_id']}: {record['source']}" for record in manifest["records"]
    )
    price_values = ",\n  ".join(
        f"({sql(row['product_id'])}, {sql(row['size'])}, {sql(row['price'])})" for row in variants
    )
    product_values = ",\n  ".join(
        f"({sql(record['product_id'])}, {sql(record['costs'][record['default_size']]['total_cost'])})"
        for record in manifest["records"]
    )
    variant_values = ",\n  ".join(
        f"({sql(row['product_id'])}, {sql(row['name'])}, {sql(row['size'])}, {sql(row['price'])}, "
        f"{sql(row['cost'])}, {sql(row['tax_rate'])}, {str(row['is_default']).lower()}, {sql(row['sort_order'])})"
        for row in variants
    )
    component_values = ",\n  ".join(
        f"({sql(row['product_id'])}, {sql(row['component_type'])}, {sql(row['name'])}, {sql(row['amount'])}, "
        f"'per_unit', date {sql(row['effective_from'])})"
        for row in components
    )
    snapshot_values = ",\n  ".join(
        f"({sql(record['product_id'])}, {sql(record['costs'][record['default_size']].get('ingredient_cost', 0))}, "
        f"{sql(record['costs'][record['default_size']].get('packaging_cost', 0))}, "
        f"{sql(record['costs'][record['default_size']].get('labor_cost', 0))}, "
        f"{sql(record['costs'][record['default_size']].get('other_cost', 0))}, "
        f"{sql(record['costs'][record['default_size']]['total_cost'])}, "
        f"{sql(record['costs'][record['default_size']]['total_cost'])}, "
        f"{sql(record['prices'][record['default_size']])}, "
        f"{sql(round(record['prices'][record['default_size']] - record['costs'][record['default_size']]['total_cost'], 2))}, "
        f"{sql(record['prices'][record['default_size']])})"
        for record in manifest["records"]
    )
    product_id_list = ", ".join(sql(pid) for pid in product_ids)

    return f"""-- ASMR & SAMR product pricing import
-- Generated from latest Excel workbook revisions on {manifest['generated_at']}.
{source_lines}
-- Public prices are ex-VAT. The storefront adds VAT at checkout.
-- IMPORTANT: This file intentionally does not adjust sellable stock.

begin;

insert into public.product_prices (product_id, size, price)
values
  {price_values}
on conflict (product_id, size) do update
set price = excluded.price;

update public.products p
set cost = v.cost,
    tax_rate = {VAT_RATE},
    availability = case when p.availability = 'discontinued' then p.availability else 'in_stock' end,
    status = case when p.status = 'archived' then p.status else 'published' end,
    updated_at = now()
from (values
  {product_values}
) as v(product_id, cost)
where p.id = v.product_id;

update public.product_variants pv
set price = v.price,
    cost = v.cost,
    tax_rate = v.tax_rate,
    is_default = v.is_default,
    sort_order = v.sort_order,
    updated_at = now()
from (values
  {variant_values}
) as v(product_id, name, size, price, cost, tax_rate, is_default, sort_order)
where pv.product_id = v.product_id
  and pv.size = v.size
  and pv.archived_at is null;

insert into public.product_variants (
  product_id, name, size, price, cost, tax_rate, is_default, is_active, sort_order
)
select v.product_id, v.name, v.size, v.price, v.cost, v.tax_rate, v.is_default, true, v.sort_order
from (values
  {variant_values}
) as v(product_id, name, size, price, cost, tax_rate, is_default, sort_order)
where not exists (
  select 1
  from public.product_variants pv
  where pv.product_id = v.product_id
    and pv.size = v.size
    and pv.archived_at is null
);

delete from public.product_cost_components
where product_id in ({product_id_list})
  and effective_from = date {sql(EFFECTIVE_DATE.isoformat())}
  and name like 'Excel launch import:%';

insert into public.product_cost_components (
  product_id, component_type, name, amount, allocation_method, effective_from
)
values
  {component_values};

insert into public.product_cost_snapshots (
  product_id, ingredient_cost, packaging_cost, labor_cost, other_cost,
  total_cost, cost_per_unit, selling_price, gross_profit, recommended_price
)
values
  {snapshot_values};

commit;
"""


def main() -> int:
    parser = argparse.ArgumentParser(description="Extract product pricing from the latest ASMR/SAMR Excel workbooks.")
    parser.add_argument("--json", default=str(DEFAULT_JSON), help="Output JSON manifest path.")
    parser.add_argument("--sql", default=str(DEFAULT_SQL), help="Output SQL import path.")
    args = parser.parse_args()

    manifest = build_manifest()
    json_path = Path(args.json)
    sql_path = Path(args.sql)
    json_path.parent.mkdir(parents=True, exist_ok=True)
    sql_path.parent.mkdir(parents=True, exist_ok=True)
    json_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    sql_path.write_text(render_sql(manifest), encoding="utf-8")
    print(f"Wrote {json_path}")
    print(f"Wrote {sql_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
